import copy
import csv
import io
import json

import pytest
from fastapi.testclient import TestClient
from PIL import Image

from agent import config, loop, results, schemas
from agent.server import app as app_module
from agent.server import catalogue, export, jobs
from agent.server.lv95 import wgs84_to_lv95

FIXTURE_ID = "epics:PK_020566"


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "CACHE_DIR", tmp_path / "cache")
    monkeypatch.setattr(jobs, "FAKE_DELAY_S", 0.001)
    monkeypatch.setattr(jobs, "_jobs", {})
    monkeypatch.delenv("AGENT_OFFLINE", raising=False)
    return TestClient(app_module.app)


@pytest.fixture
def fake(monkeypatch):
    monkeypatch.setattr(config, "FAKE", True)


@pytest.fixture
def real(monkeypatch):
    monkeypatch.setattr(config, "FAKE", False)


def fixture_result(photo_id=FIXTURE_ID):
    r = json.loads((config.FIXTURES / "analysis_result.json").read_text(encoding="utf-8"))
    r["photo_id"] = photo_id
    return r


def sse_events(client, job_id):
    events, event = [], None
    with client.stream("GET", f"/analyze/{job_id}/events") as resp:
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("text/event-stream")
        for line in resp.iter_lines():
            if line.startswith("event: "):
                event = line[7:]
            elif line.startswith("data: "):
                events.append((event, json.loads(line[6:])))
                if event == "result":
                    break
    return events


def test_health(client):
    r = client.get("/health").json()
    assert r["ok"] is True and r["model"] == config.MODEL
    assert r["offline"] is False and r["tiles_url"].startswith("https://wmts.geo.admin.ch")


def test_health_offline(client, monkeypatch):
    monkeypatch.setenv("AGENT_OFFLINE", "1")
    r = client.get("/health").json()
    assert r["offline"] is True and r["tiles_url"] == "/tiles"


def test_cors_localhost(client):
    r = client.get("/health", headers={"Origin": "http://localhost:3000"})
    assert r.headers["access-control-allow-origin"] == "http://localhost:3000"


def test_photos_stub(client):
    r = client.get("/photos").json()
    assert r["total"] == 1 and r["items"][0]["id"] == FIXTURE_ID
    schemas.validate(r["items"][0], "photo_record")
    assert client.get("/photos", params={"q": "urirotstock"}).json()["total"] == 1
    assert client.get("/photos", params={"municipality": "Schwyz"}).json()["total"] == 0
    assert client.get("/photos", params={"split": "reference"}).json()["total"] == 1
    assert client.get("/photos", params={"offset": 1}).json()["items"] == []
    assert client.get(f"/photos/{FIXTURE_ID}").json()["id"] == FIXTURE_ID
    assert client.get("/photos/stasz:999").status_code == 404


def test_images_stub(client):
    r = client.get(f"/images/{FIXTURE_ID}")
    assert r.status_code == 200 and r.headers["content-type"] == "image/jpeg"
    full = Image.open(io.BytesIO(r.content))
    r = client.get(f"/images/{FIXTURE_ID}", params={"w": 200})
    small = Image.open(io.BytesIO(r.content))
    assert small.width == 200 and full.width > 200
    assert client.get("/images/stasz:999").status_code == 404


def test_fake_analyze_sse(client, fake):
    r = client.post("/analyze", json={"photo_id": "stasz:123"})
    assert r.status_code == 200
    job_id = r.json()["job_id"]
    assert r.json()["photo_id"] == "stasz:123"
    events = sse_events(client, job_id)
    fx = fixture_result()
    assert [e for e, _ in events] == ["trace"] * len(fx["trace"]) + ["result"]
    assert [d for e, d in events if e == "trace"] == fx["trace"]
    final = events[-1][1]
    schemas.validate(final, "analysis_result")
    assert final["photo_id"] == "stasz:123" and final["status"] == "done"
    # a late client gets everything again
    assert len(sse_events(client, job_id)) == len(events)
    assert client.get(f"/analyze/{job_id}").json()["status"] == "done"
    # fake results stay in memory only, but /results serves them
    assert results.load("stasz:123") is None
    assert client.get("/results/stasz:123").json()["photo_id"] == "stasz:123"


def test_running_job_status(client, fake, monkeypatch):
    monkeypatch.setattr(jobs, "FAKE_DELAY_S", 0.2)
    job_id = client.post("/analyze", json={"photo_id": "stasz:5"}).json()["job_id"]
    snap = client.get(f"/analyze/{job_id}").json()
    assert snap["status"] == "running" and snap["photo_id"] == "stasz:5"
    assert client.get("/analyze/nope").status_code == 404


def test_cache_hit_and_fresh(client, real, monkeypatch):
    calls = []

    def fake_analyze(photo_id=None, image_path=None, *, model=config.MODEL, exclude_ids=(), on_event=None):
        calls.append(photo_id)
        r = fixture_result(photo_id)
        r["run_id"] = f"run-{len(calls)}"
        for item in r["trace"]:
            on_event(item)
        return r

    monkeypatch.setattr(loop, "analyze", fake_analyze)
    cached = fixture_result()
    cached["run_id"] = "cached"
    results.save(cached)

    r = client.post("/analyze", json={"photo_id": FIXTURE_ID}).json()
    snap = client.get(f"/analyze/{r['job_id']}").json()
    assert snap["status"] == "done" and snap["run_id"] == "cached" and calls == []
    assert sse_events(client, r["job_id"])[-1][1]["run_id"] == "cached"

    r = client.post("/analyze", params={"fresh": 1}, json={"photo_id": FIXTURE_ID}).json()
    final = sse_events(client, r["job_id"])[-1][1]
    assert calls == [FIXTURE_ID] and final["run_id"] == "run-1"
    assert results.load(FIXTURE_ID)["run_id"] == "run-1"


def test_real_analyze_unknown_photo(client, real):
    assert client.post("/analyze", json={"photo_id": "stasz:999"}).status_code == 404
    assert client.post("/analyze", json={}).status_code == 400


def test_real_analyze_error_ends_stream(client, real, monkeypatch):
    def boom(*a, **kw):
        raise RuntimeError("kaputt")

    monkeypatch.setattr(loop, "analyze", boom)
    job_id = client.post("/analyze", json={"photo_id": FIXTURE_ID}).json()["job_id"]
    final = sse_events(client, job_id)[-1]
    assert final[0] == "result" and final[1]["status"] == "error"
    assert final[1]["error"] == "Die Analyse ist fehlgeschlagen." and "kaputt" in final[1]["error_detail"]


def test_missing_key_gives_german_error(client, real, monkeypatch):
    from agent import llm

    def no_key(*a, **kw):
        raise llm.MissingKey("ANTHROPIC_API_KEY is not set")

    monkeypatch.setattr(loop, "analyze", no_key)
    job_id = client.post("/analyze", json={"photo_id": FIXTURE_ID}).json()["job_id"]
    final = sse_events(client, job_id)[-1]
    assert final[1]["error"].startswith("Kein API-Schlüssel")


def test_upload(client, fake):
    buf = io.BytesIO()
    Image.new("RGB", (1600, 800), (120, 130, 140)).save(buf, "PNG")
    r = client.post("/analyze", files={"file": ("karte.png", buf.getvalue(), "image/png")}).json()
    pid = r["photo_id"]
    assert pid.startswith("upload:") and len(pid) == len("upload:") + 12
    assert sse_events(client, r["job_id"])[-1][1]["photo_id"] == pid
    rec = client.get(f"/photos/{pid}").json()
    schemas.validate(rec, "photo_record")
    img = Image.open(io.BytesIO(client.get(f"/images/{pid}").content))
    assert max(img.size) == 1024
    bad = client.post("/analyze", files={"file": ("x.png", b"not an image", "image/png")})
    assert bad.status_code == 400


def test_results_404(client):
    assert client.get("/results/stasz:404").status_code == 404


def _csv_rows(content: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(content.decode("utf-8-sig")), delimiter=";"))


def test_export_csv(client):
    results.save(fixture_result())
    r = client.get("/export.csv")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert r.content.startswith(b"\xef\xbb\xbf")
    got, exp = _csv_rows(r.content), _csv_rows((config.FIXTURES / "export.csv").read_bytes())
    assert got[0] == exp[0] == export.COLUMNS
    row, want = dict(zip(got[0], got[1])), dict(zip(exp[0], exp[1]))
    # same values as the contract fixture, except the columns the fixture fills by hand
    for col in export.COLUMNS:
        if col not in ("Aufnahmedatum", "Quellen"):  # record.date.raw is "vor 13.6.1907"; Quellen picks differ
            assert row[col] == want[col], col
    assert row["Nicht_bestimmbar"] == "Strasse" and row["Pruefung_noetig"] == "ja" and row["Konfidenz_Gemeinde"] == "0.82"


def test_export_json_and_ids(client):
    results.save(fixture_result())
    results.save(fixture_result("stasz:2"))
    rows = client.get("/export.json").json()
    assert len(rows) == 2 and set(rows[0]) == set(export.COLUMNS)
    rows = client.get("/export.json", params={"ids": f"{FIXTURE_ID},stasz:missing"}).json()
    assert len(rows) == 1 and rows[0]["E_LV95"] == 2689007 and rows[0]["Dateiname"] == "PK_020566.jpg"
    lines = _csv_rows(client.get("/export.csv", params={"ids": "stasz:2"}).content)
    assert len(lines) == 2 and lines[1][0] == "stasz:2"  # no record for stasz:2 in stub mode: id as key


def test_export_xlsx_has_both_sheets_and_links(client):
    from openpyxl import load_workbook

    results.save(fixture_result())
    r = client.get("/export.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Metadaten", "Methode"]
    ws = wb["Metadaten"]
    assert [c.value for c in ws[1]] == export.COLUMNS
    assert ws["A2"].value == "PK_020566.jpg" and ws["A2"].hyperlink.target.endswith("/images/epics:PK_020566")
    gemeinde = ws.cell(row=2, column=export.COLUMNS.index("Gemeinde") + 1).value
    assert gemeinde == "Ingenbohl"
    method = {row[0].value: row[1].value for row in wb["Methode"].iter_rows(min_row=2)}
    from agent import confidence
    assert method["3 Quellen"] == f"{confidence.SCALE[3]:.2f}" and method["4 oder mehr Quellen"] == f"{confidence.TOP:.2f}"


def test_review_confirms_and_corrects(client):
    results.save(fixture_result())
    assert client.post(f"/review/{FIXTURE_ID}", json={"status": "egal"}).status_code == 400
    assert client.post(f"/review/{FIXTURE_ID}", json={"status": "korrigiert"}).status_code == 400
    assert client.post(f"/review/{FIXTURE_ID}", json={"status": "bestaetigt", "fields": {"decade": 1900}}).status_code == 400
    assert client.post("/review/stasz:missing", json={"status": "bestaetigt"}).status_code == 404
    r = client.post(f"/review/{FIXTURE_ID}", json={"status": "korrigiert", "fields": {"street": "Axenstrasse 2"},
                                                     "note": "vor Ort geprüft"}).json()
    assert r["review"]["status"] == "korrigiert" and r["review"]["reviewed_at"]
    assert results.load(FIXTURE_ID)["review"]["fields"] == {"street": "Axenstrasse 2"}
    row = client.get("/export.json").json()[0]
    assert row["Strasse"] == "Axenstrasse 2" and row["Review_Status"] == "korrigiert"
    assert row["Pruefung_noetig"] == "nein" and row["Nicht_bestimmbar"] is None


def test_photos_priority_filter_and_order(monkeypatch):
    def rec(i, prio):
        return {"id": f"challenge:SG.CIV.17.{i}", "signature": f"SG.CIV.17.{i}", "split": "challenge",
                "raw": {"priority": prio}}
    recs = [rec(10, 2), rec(3, 1), rec(25, 1), rec(2, 2)]
    monkeypatch.setattr(catalogue, "all_records", lambda: recs)
    ids = [r["id"] for r in catalogue.list_records("challenge")]
    assert ids == ["challenge:SG.CIV.17.3", "challenge:SG.CIV.17.25", "challenge:SG.CIV.17.2", "challenge:SG.CIV.17.10"]
    assert [r["id"] for r in catalogue.list_records("challenge", priority=1)] == ids[:2]


def test_lv95():
    e, n = wgs84_to_lv95(46.951082877, 7.438632495)  # swisstopo reference point (Bern)
    assert abs(e - 2600000) < 2 and abs(n - 1200000) < 2


def test_timeline(client):
    r = client.get("/timeline", params={"lat": 46.98, "lon": 8.6}).json()
    assert r["photos"][0]["image_url"] == f"/images/{FIXTURE_ID}"
    assert set(r["photos"][0]) == {"id", "year_from", "year_to", "distance_m", "image_url"}
    assert r["map_years"] and all(isinstance(y, int) for y in r["map_years"])


def test_eval_summary_fixture_then_real(client):
    fx = json.loads((config.FIXTURES / "eval_summary.json").read_text(encoding="utf-8"))
    assert client.get("/eval/summary").json() == fx
    config.RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    (config.RESULTS_DIR / "eval_summary.json").write_text(json.dumps({**fx, "n": 7}), encoding="utf-8")
    assert client.get("/eval/summary").json()["n"] == 7


def test_demo_stub(client):
    r = client.get("/demo").json()
    assert r and r[0]["id"] == FIXTURE_ID and r[0]["record"]["id"] == FIXTURE_ID and r[0]["why"]


def test_tiles_cached_and_offline(client, monkeypatch):
    monkeypatch.setenv("AGENT_OFFLINE", "1")
    p = config.CACHE_DIR / "tiles" / "ch.swisstopo.zeitreihen" / "19041231" / "14" / "8580" / "5772.png"
    p.parent.mkdir(parents=True)
    p.write_bytes(b"\x89PNG fake")
    r = client.get("/tiles/ch.swisstopo.zeitreihen/19041231/14/8580/5772")
    assert r.status_code == 200 and r.headers["content-type"] == "image/png" and r.content == b"\x89PNG fake"
    assert client.get("/tiles/ch.swisstopo.zeitreihen/19041231/14/8580/5773").status_code == 404
    assert client.get("/tiles/..%2Fetc/x/1/2/3").status_code in (400, 404)
